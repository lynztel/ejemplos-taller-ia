from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pathlib import Path

OUT = Path(r"C:\Users\Admin\Desktop\MiBiblioteca\01 Projects\PISA México - Resultados e impacto económico\02 Presentación\PISA_Mexico_impacto_economico.pptx")
TMP = OUT.parent / "_charts"
TMP.mkdir(exist_ok=True)

NAVY = RGBColor(18, 43, 70)
BLUE = RGBColor(35, 105, 170)
TEAL = RGBColor(35, 145, 145)
ORANGE = RGBColor(226, 123, 48)
RED = RGBColor(190, 67, 67)
INK = RGBColor(35, 42, 50)
MUTED = RGBColor(95, 106, 117)
LIGHT = RGBColor(244, 247, 250)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(57, 130, 88)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

def add_bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid(); fill.fore_color.rgb = color

def textbox(slide, x, y, w, h, text, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.margin_left = Inches(.04); tf.margin_right = Inches(.04)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text; r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return box

def richbox(slide, x, y, w, h, runs, size=18, align=PP_ALIGN.LEFT, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    for txt, col, bold in runs:
        r = p.add_run(); r.text = txt; r.font.name = "Aptos"; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = col
    return box

def title(slide, kicker, heading, subtitle=None):
    textbox(slide, .65, .35, 12, .28, kicker.upper(), 10, BLUE, True)
    textbox(slide, .65, .68, 12, .62, heading, 27, NAVY, True)
    if subtitle: textbox(slide, .68, 1.34, 11.8, .35, subtitle, 12, MUTED)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.65), Inches(1.82), Inches(1.1), Inches(.05))
    line.fill.solid(); line.fill.fore_color.rgb = ORANGE; line.line.fill.background()

def footer(slide, source="Fuente: OECD, PISA 2025 Results (Volume I), Country Note: Mexico. Consulta: 10 oct 2026."):
    textbox(slide, .65, 7.12, 12, .18, source, 7.5, MUTED)
    textbox(slide, 12.55, 7.08, .25, .22, str(len(prs.slides)), 8, MUTED, False, PP_ALIGN.RIGHT)

def card(slide, x, y, w, h, heading, body, accent=BLUE, body_size=14):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = LIGHT; sh.line.color.rgb = RGBColor(222, 229, 236)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(.08), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    textbox(slide, x+.23, y+.22, w-.42, .35, heading, 15, NAVY, True)
    textbox(slide, x+.23, y+.7, w-.42, h-.85, body, body_size, INK)

def add_chart(slide, chart_type, x, y, w, h, categories, series, colors, max_scale=None, show_legend=False):
    data = CategoryChartData(); data.categories = categories
    for name, vals in series: data.add_series(name, vals)
    chart = slide.shapes.add_chart(chart_type, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    chart.has_legend = show_legend
    if show_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM; chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(9)
    chart.value_axis.tick_labels.font.size = Pt(9); chart.value_axis.tick_labels.font.color.rgb = MUTED
    chart.category_axis.tick_labels.font.size = Pt(9); chart.category_axis.tick_labels.font.color.rgb = MUTED
    chart.value_axis.has_major_gridlines = True
    if max_scale: chart.value_axis.maximum_scale = max_scale
    for i, s in enumerate(chart.series):
        s.format.fill.solid(); s.format.fill.fore_color.rgb = colors[i]
        s.format.line.color.rgb = colors[i]
    return chart

def transition(slide):
    # Fade transition: subtle academic animation when advancing slides.
    sld = slide._element
    trans = parse_xml(f'<p:transition {nsdecls("p")} spd="med" advClick="1"><p:fade/></p:transition>')
    sld.insert(2, trans)

# 1
s = prs.slides.add_slide(blank); add_bg(s, NAVY)
textbox(s, .8, .75, 11.8, .35, "PISA 2025 · MÉXICO", 13, RGBColor(155, 205, 226), True)
textbox(s, .8, 1.45, 11.2, 1.45, "¿Qué precio económico\ntiene aprender menos?", 34, WHITE, True)
textbox(s, .85, 3.35, 10.5, .62, "Resultados, interpretación y repercusiones a corto y mediano plazo", 18, RGBColor(224, 232, 240))
textbox(s, .85, 5.72, 5.8, .38, "Presentación académica · 10 de octubre de 2026", 11, RGBColor(180, 197, 212))
accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.7), Inches(1.0), Inches(1.4), Inches(5.2)); accent.fill.solid(); accent.fill.fore_color.rgb = ORANGE; accent.line.fill.background(); accent.rotation = 12
textbox(s, 10.75, 2.2, 1.3, 1.8, "388\nMAT", 24, WHITE, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
textbox(s, 10.75, 4.2, 1.3, 1.2, "México\n2025", 14, WHITE, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)

# 2
s = prs.slides.add_slide(blank); add_bg(s); title(s, "Resumen ejecutivo", "Una señal educativa con consecuencias productivas", "PISA no mide el PIB; mide capacidades que condicionan el potencial de productividad futura.")
card(s, .7, 2.25, 3.0, 2.3, "388 puntos", "México en matemáticas; 75 puntos por debajo del promedio OCDE.", ORANGE, 16)
card(s, 3.95, 2.25, 3.0, 2.3, "50%", "Alcanza al menos el nivel base en lectura y ciencias; en matemáticas, 30%.", BLUE, 16)
card(s, 7.2, 2.25, 3.0, 2.3, "68 puntos", "Brecha en ciencias entre el 25% más aventajado y el 25% más desfavorecido.", RED, 16)
card(s, 10.45, 2.25, 2.15, 2.3, "Horizonte", "Corto plazo: capacitación y empleabilidad.\n\nMediano plazo: innovación, productividad y movilidad.", TEAL, 12)
textbox(s, .85, 5.05, 11.6, .85, "Tesis: la prioridad económica no es subir un ranking; es ampliar la base de jóvenes capaces de leer, modelar, resolver problemas y aprender durante toda la vida.", 20, NAVY, True, PP_ALIGN.CENTER)
footer(s)

# 3
s = prs.slides.add_slide(blank); add_bg(s); title(s, "Marco de lectura", "¿Qué mide PISA y qué no mide?", "La edición 2025 evaluó a estudiantes de 15 años en 91 países y economías.")
card(s, .8, 2.25, 3.65, 2.25, "Ciencias", "Explicar fenómenos científicamente, evaluar evidencia y diseñar investigaciones.", TEAL, 15)
card(s, 4.85, 2.25, 3.65, 2.25, "Matemáticas", "Interpretar situaciones, modelar problemas y elegir estrategias.", ORANGE, 15)
card(s, 8.9, 2.25, 3.65, 2.25, "Lectura", "Comprender, evaluar y usar textos para aprender y decidir.", BLUE, 15)
textbox(s, 1.0, 5.05, 5.35, .9, "PISA sí permite", 16, NAVY, True)
textbox(s, 1.0, 5.45, 5.35, .8, "Comparar competencias, observar tendencias y detectar brechas de equidad.", 15, INK)
textbox(s, 7.0, 5.05, 5.0, .9, "PISA no permite", 16, NAVY, True)
textbox(s, 7.0, 5.45, 5.0, .8, "Atribuir por sí sola una causa política o calcular directamente el impacto en el PIB.", 15, INK)
footer(s, "Fuente: OECD (2026), PISA 2025 Results, Country Note: Mexico.")

# 4
s = prs.slides.add_slide(blank); add_bg(s); title(s, "Resultados 2025", "México se mantiene por debajo del promedio OCDE", "La brecha es mayor en ciencias y matemáticas; los tres dominios están por debajo del promedio internacional.")
add_chart(s, XL_CHART_TYPE.BAR_CLUSTERED, .85, 2.15, 7.0, 4.2, ["Matemáticas", "Lectura", "Ciencias"], [("México", [388,408,414]), ("OCDE", [463,461,482])], [ORANGE, BLUE], 520, True)
card(s, 8.35, 2.35, 3.9, 1.2, "−75 puntos", "Brecha México–OCDE en matemáticas.", ORANGE, 14)
card(s, 8.35, 3.8, 3.9, 1.2, "−68 puntos", "Brecha México–OCDE en ciencias.", TEAL, 14)
card(s, 8.35, 5.25, 3.9, 1.2, "−53 puntos", "Brecha México–OCDE en lectura.", BLUE, 14)
footer(s, "Fuente: OECD (2026), PISA 2025 Results; Education GPS, Mexico.")

# 5
s = prs.slides.add_slide(blank); add_bg(s); title(s, "Tendencia", "El reto no comienza en 2025", "Matemáticas y lectura retroceden frente a 2022; ciencias permanece aproximadamente estable.")
add_chart(s, XL_CHART_TYPE.LINE_MARKERS, .75, 2.1, 7.3, 4.45, ["2015", "2018", "2022", "2025"], [("Matemáticas", [408,409,395,388]), ("Lectura", [423,420,415,408]), ("Ciencias", [416,419,410,414])], [ORANGE, BLUE, TEAL], 450, True)
card(s, 8.45, 2.35, 3.7, 1.4, "Matemáticas", "−7 puntos vs. 2022; el nivel base alcanza solo a 30%.", ORANGE, 14)
card(s, 8.45, 4.0, 3.7, 1.4, "Lectura", "−7 puntos vs. 2022; 50% llega al nivel base.", BLUE, 14)
card(s, 8.45, 5.65, 3.7, .9, "Ciencias", "+4 puntos vs. 2022; la OCDE lo describe como estable.", TEAL, 12)
footer(s, "Fuente: OECD, PISA 2025 Country Note: Mexico; serie histórica PISA 2015–2025.")

# 6
s = prs.slides.add_slide(blank); add_bg(s); title(s, "Competencias mínimas", "La base de habilidades es estrecha", "Porcentaje de estudiantes que alcanza al menos el nivel 2 de desempeño.")
add_chart(s, XL_CHART_TYPE.BAR_CLUSTERED, .9, 2.15, 7.2, 4.1, ["Matemáticas", "Lectura", "Ciencias"], [("México", [30,50,50]), ("OCDE", [65,69,74])], [ORANGE, BLUE], 80, True)
textbox(s, 8.55, 2.45, 3.65, .8, "Nivel 2 = umbral de competencia funcional", 17, NAVY, True)
textbox(s, 8.55, 3.35, 3.65, 2.3, "Debajo de este umbral, es más difícil interpretar información, tomar decisiones cuantitativas y aprender de manera autónoma en contextos nuevos.", 16, INK)
textbox(s, 8.55, 5.95, 3.65, .6, "Casi no hay estudiantes mexicanos en niveles 5–6.", 14, RED, True)
footer(s)

# 7
s = prs.slides.add_slide(blank); add_bg(s); title(s, "Equidad", "El promedio nacional oculta una brecha social", "La desventaja socioeconómica sigue asociada con menores resultados, aunque la brecha mexicana es menor que el promedio OCDE.")
card(s, .9, 2.25, 3.4, 2.0, "68 puntos", "Diferencia en ciencias entre el 25% más aventajado y el 25% más desfavorecido.", RED, 16)
card(s, 4.6, 2.25, 3.4, 2.0, "12%", "Proporción de la variación en ciencias asociada al estatus socioeconómico.", ORANGE, 16)
card(s, 8.3, 2.25, 3.4, 2.0, "12%", "De estudiantes desfavorecidos alcanzó el cuartil superior de desempeño en ciencias: resiliencia académica.", TEAL, 14)
textbox(s, 1.0, 5.0, 11.2, .7, "Implicación económica: si el origen predice demasiado el aprendizaje, se reduce la movilidad social y se desaprovecha talento productivo.", 20, NAVY, True, PP_ALIGN.CENTER)
textbox(s, 1.2, 5.95, 10.9, .55, "La expansión de la cobertura puede incorporar a más jóvenes vulnerables; eso puede hacer que el promedio parezca más bajo sin que toda la caída sea pérdida de aprendizaje.", 12, MUTED, False, PP_ALIGN.CENTER)
footer(s)

# 8
s = prs.slides.add_slide(blank); add_bg(s); title(s, "Interpretación económica", "De la habilidad escolar al desempeño productivo", "El vínculo es un mecanismo plausible, no una causalidad directa demostrada por PISA.")
steps = [("Aprendizaje", "lectura · matemática · ciencia", BLUE), ("Habilidades", "resolver · comunicar · aprender", TEAL), ("Mercado laboral", "empleabilidad · capacitación", ORANGE), ("Economía", "productividad · innovación · ingresos", GREEN)]
for i, (h,b,c) in enumerate(steps):
    x = .75 + i*3.12
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.55), Inches(2.6), Inches(1.45)); sh.fill.solid(); sh.fill.fore_color.rgb = LIGHT; sh.line.color.rgb = c
    textbox(s, x+.12, 2.8, 2.35, .35, h, 17, c, True, PP_ALIGN.CENTER)
    textbox(s, x+.15, 3.28, 2.3, .45, b, 12, INK, False, PP_ALIGN.CENTER)
    if i < 3:
        textbox(s, x+2.68, 2.9, .45, .4, "→", 24, ORANGE, True, PP_ALIGN.CENTER)
card(s, 1.05, 4.75, 5.3, 1.25, "Corto plazo · 1–5 años", "Más gasto de nivelación y capacitación; transición escuela-trabajo menos eficiente.", ORANGE, 14)
card(s, 6.95, 4.75, 5.3, 1.25, "Mediano plazo · 5–15 años", "Menor adopción tecnológica y menor oferta de talento para sectores sofisticados si la brecha persiste.", BLUE, 14)
footer(s, "Interpretación: síntesis económica basada en mecanismos de capital humano; no es una estimación causal de PISA.")

# 9
s = prs.slides.add_slide(blank); add_bg(s); title(s, "Escenarios", "El costo de no actuar es acumulativo", "Índice ilustrativo de capacidad productiva futura — no es una proyección del PIB.")
add_chart(s, XL_CHART_TYPE.LINE_MARKERS, .85, 2.1, 7.3, 4.4, ["2025", "2030", "2035", "2040"], [("Continuidad", [100,101,101,102]), ("Mejora gradual", [100,104,109,114])], [RED, GREEN], 120, True)
card(s, 8.45, 2.4, 3.7, 1.45, "Continuidad", "Persisten brechas básicas, costos de capacitación y menor movilidad.", RED, 14)
card(s, 8.45, 4.1, 3.7, 1.45, "Mejora gradual", "Fortalecer fundamentos, docentes y transición digital eleva el potencial, con rezagos iniciales.", GREEN, 14)
textbox(s, 8.45, 5.9, 3.7, .55, "La prioridad es cambiar la trayectoria, no perseguir un ranking.", 14, NAVY, True, PP_ALIGN.CENTER)
footer(s, "Nota: escenarios conceptuales para comunicar mecanismos; no representan una predicción estadística.")

# 10
s = prs.slides.add_slide(blank); add_bg(s); title(s, "Agenda de respuesta", "Tres palancas para convertir diagnóstico en inversión", "Una respuesta económica debe combinar calidad, equidad y conexión con el trabajo.")
card(s, .8, 2.25, 3.75, 2.55, "1 · Fundamentos", "Priorizar lectura, razonamiento matemático y ciencia desde los primeros años.\n\nMedir aprendizaje real, no solo asistencia.", BLUE, 14)
card(s, 4.8, 2.25, 3.75, 2.55, "2 · Docentes y escuelas", "Formación continua, acompañamiento pedagógico y materiales/infrastructura focalizados en las escuelas con mayor rezago.", ORANGE, 14)
card(s, 8.8, 2.25, 3.75, 2.55, "3 · Escuela–economía", "Habilidades digitales, resolución de problemas, orientación vocacional y vínculos con empresas y educación superior.", TEAL, 14)
textbox(s, 1.0, 5.45, 11.4, .8, "Indicadores de seguimiento: nivel base, abandono, brechas regionales, inserción laboral, capacitación empresarial y productividad sectorial.", 17, NAVY, True, PP_ALIGN.CENTER)
footer(s, "Recomendaciones: interpretación propia a partir de PISA 2025 y mecanismos de capital humano.")

# 11
s = prs.slides.add_slide(blank); add_bg(s, NAVY)
textbox(s, .8, .55, 11.7, .3, "CIERRE", 11, RGBColor(155, 205, 226), True)
textbox(s, .8, 1.2, 11.4, .75, "La educación es infraestructura económica", 29, WHITE, True)
textbox(s, .85, 2.25, 11.0, 1.0, "México necesita ampliar la base de competencias fundamentales y convertir la resiliencia de sus estudiantes en productividad, innovación y movilidad social.", 22, RGBColor(224,232,240), True)
for i,(h,b,c) in enumerate([("Priorizar","fundamentos",BLUE),("Focalizar","donde la brecha es mayor",ORANGE),("Evaluar","resultados y efectos",TEAL)]):
    x=.95+i*4.05; sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(4.25),Inches(3.35),Inches(1.2)); sh.fill.solid(); sh.fill.fore_color.rgb=c; sh.line.fill.background(); textbox(s,x+.1,4.45,3.15,.3,h,17,WHITE,True,PP_ALIGN.CENTER); textbox(s,x+.1,4.82,3.15,.3,b,13,WHITE,False,PP_ALIGN.CENTER)
textbox(s, .85, 6.45, 11.6, .35, "PISA muestra la señal; las políticas públicas deben construir la trayectoria.", 14, RGBColor(180,197,212), False, PP_ALIGN.CENTER)

# 12
s = prs.slides.add_slide(blank); add_bg(s); title(s, "Referencias", "Fuentes principales", "Consulta: 10 de octubre de 2026.")
refs = [
    "OECD (2026). PISA 2025 Results (Volume I): Mexico — Country Note.",
    "OECD (2026). Education GPS — Mexico: Student performance (PISA 2025).",
    "OECD (2023). PISA 2022 Results (Volume I and II): Mexico — Country Note.",
    "OECD (2026). PISA 2025 Results (Volume I): Future-Ready Students.",
]
for i, r in enumerate(refs):
    textbox(s, 1.0, 2.2+i*.7, 11.2, .45, f"{i+1}. {r}", 16, INK, i==0)
textbox(s, 1.0, 5.55, 11.1, .75, "Enlaces:\nwww.oecd.org/en/publications/pisa-2025-results-volume-i-country-notes_2d4ff9ea-en/mexico_85702d30-en.html\ngpseducation.oecd.org/CountryProfile?primaryCountry=MEX&topic=PI&treshold=10", 11, BLUE)
footer(s, "Elaboración propia con datos OECD/PISA; las interpretaciones económicas están señaladas como mecanismos o escenarios.")

for slide in prs.slides: transition(slide)
prs.save(OUT)
print(OUT)
